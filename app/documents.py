from dataclasses import dataclass
import hashlib
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader

from .errors import UnsupportedFormatError

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}


@dataclass(frozen=True)
class ParsedDocument:
    document_id: str
    filename: str
    text: str
    chunks: list[str]


def parse_file(path: Path, chunk_size: int, chunk_overlap: int) -> ParsedDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"不支持的文件格式：{suffix or '无扩展名'}。请上传 PDF、DOCX、Markdown 或 TXT。"
        )
    text = _read_text(path, suffix).strip()
    if not text:
        raise ValueError(f"文档“{path.name}”没有可提取的文本内容。")
    document_id = hashlib.sha256(path.read_bytes()).hexdigest()
    return ParsedDocument(
        document_id=document_id,
        filename=path.name,
        text=text,
        chunks=split_text(text, chunk_size, chunk_overlap),
    )


def _read_text(path: Path, suffix: str) -> str:
    if suffix == ".pdf":
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if suffix == ".docx":
        return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)
    return path.read_text(encoding="utf-8-sig", errors="replace")


def split_text(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
    if chunk_overlap >= chunk_size:
        raise ValueError("CHUNK_OVERLAP 必须小于 CHUNK_SIZE。")
    normalized = re.sub(r"[ \t]+\n", "\n", text).strip()
    chunks: list[str] = []
    start = 0
    while start < len(normalized):
        end = min(len(normalized), start + chunk_size)
        if end < len(normalized):
            boundary = max(normalized.rfind("\n", start, end), normalized.rfind(" ", start, end))
            if boundary > start + chunk_size // 2:
                end = boundary
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            break
        start = max(start + 1, end - chunk_overlap)
    return chunks

